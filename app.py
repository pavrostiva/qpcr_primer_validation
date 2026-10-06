import os
import glob
import re
import io
import tempfile
import zipfile
import numpy as np
import pandas as pd
from scipy import stats
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pdf_report import build_pdf_report

st.set_page_config(
    page_title="qPCR primer validation",
    page_icon="🧬",
    layout="wide"
)

# ----------------- Helper & Parser Functions -----------------

def normalize_well(val: str) -> str:
    """Normalize well identifiers to standard A01-H12 format."""
    m = re.search(r"([A-Ha-h])\s*0?(\d{1,2})", str(val))
    return f"{m.group(1).upper()}{int(m.group(2)):02d}" if m else str(val).strip()


def parse_dilution_value(val_str: str) -> float:
    """Parse biological dilutions, supporting '1k' (1000) and '10k' (10000)."""
    s = str(val_str).lower().strip()
    if s.endswith("k"):
        return float(s[:-1]) * 1000.0
    return float(s)


def format_dilution(val) -> str:
    """Safely format dilution factors into clean strings without NaN crashes."""
    if pd.isna(val) or val is None:
        return "Unknown"
    try:
        f_val = float(val)
        if f_val.is_integer():
            return f"1/{int(f_val)}"
        return f"1/{f_val:.1f}"
    except (ValueError, TypeError):
        return str(val)


def find_files_recursively(base_dir: str):
    """
    Search recursively for Bio-Rad files and user plate layouts.
    Strictly blacklists thermocycler system files from layout matching.
    """
    files_map = {"cq": None, "amp": None, "melt": None, "layout": None}
    layout_candidates = []

    BIORAD_SYSTEM_KEYWORDS = [
        "quantification", "melt curve", "end point", "results_sybr", 
        "run information", "summary", "standard curve results"
    ]

    for root, dirs, files in os.walk(base_dir):
        dirs[:] = [d for d in dirs if d != "__MACOSX"]
        for f in files:
            if f.startswith(("._", "~$")):
                continue  # macOS resource forks and Excel lock files
            full_path = os.path.join(root, f)
            f_lower = f.lower()

            if "quantification cq results" in f_lower and f.endswith(".csv"):
                files_map["cq"] = full_path
            elif "quantification amplification results" in f_lower and f.endswith(".csv"):
                files_map["amp"] = full_path
            elif "melt curve derivative results" in f_lower and f.endswith(".csv"):
                files_map["melt"] = full_path

            if any(k in f_lower for k in BIORAD_SYSTEM_KEYWORDS):
                continue

            if f_lower.endswith((".xlsx", ".xls", ".csv")):
                score = 0
                name_no_ext = os.path.splitext(f_lower)[0]
                if name_no_ext in ["plate", "layout", "plate_layout", "platemap", "plate_map", "plan", "grid"]:
                    score = 100 if f_lower.endswith((".xlsx", ".xls")) else 80
                elif any(k in name_no_ext for k in ["plate", "layout", "plan", "grid"]):
                    score = 50 if f_lower.endswith((".xlsx", ".xls")) else 30

                if score > 0:
                    layout_candidates.append((score, full_path))

    if layout_candidates:
        layout_candidates.sort(key=lambda x: x[0], reverse=True)
        files_map["layout"] = layout_candidates[0][1]

    return files_map


def parse_plate_file(layout_filepath: str) -> pd.DataFrame:
    """
    Parse 8x12 plate layout grid from Excel or CSV.
    Safely handles empty/unpipetted wells without column shifting or KeyError.
    """
    if layout_filepath.endswith((".xlsx", ".xls")):
        df_raw = pd.read_excel(layout_filepath, header=None)
    else:
        df_raw = pd.read_csv(layout_filepath, header=None)

    df_raw = df_raw.fillna("").astype(str)
    records = []
    rows_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
    COPY_PASTE_10X_MAP = {"1": 1.0, "4": 10.0, "16": 100.0, "64": 1000.0, "256": 10000.0}

    for r_letter in rows_letters:
        row_match = df_raw[df_raw.iloc[:, 0].str.strip().str.upper() == r_letter]
        start_col_offset = 1
        if row_match.empty and df_raw.shape[1] > 1:
            row_match = df_raw[df_raw.iloc[:, 1].str.strip().str.upper() == r_letter]
            start_col_offset = 2

        row_idx = row_match.index[0] if not row_match.empty else None

        for col_idx in range(1, 13):
            well_id = f"{r_letter}{col_idx:02d}"
            tag = ""

            if row_idx is not None:
                actual_col = start_col_offset + (col_idx - 1)
                if actual_col < df_raw.shape[1]:
                    tag = df_raw.iloc[row_idx, actual_col].strip()

            # Empty unpipetted well
            if not tag or tag.lower() in ["empty", "unused", "none", "nan", "-"]:
                records.append({
                    "Well": well_id,
                    "Row": r_letter,
                    "Col": col_idx,
                    "Primer": "Empty",
                    "Series": "None",
                    "Tag": "Empty",
                    "Dilution_Factor": np.nan,
                    "Is_NTC": False,
                    "Is_Empty": True
                })
                continue

            parts = [p.strip() for p in tag.split("_") if p.strip()]

            if len(parts) >= 3:
                primer, series, d_str = "_".join(parts[:-2]), parts[-2], parts[-1]
            elif len(parts) == 2:
                primer, series, d_str = parts[0], "std", parts[1]
            else:
                primer, series, d_str = parts[0] if parts else tag, "std", "1"

            is_ntc = d_str.upper() == "NTC"
            factor = np.nan

            if not is_ntc:
                clean_s = series.lower().strip()
                clean_d = d_str.lower().strip()
                if "10" in clean_s and clean_d in COPY_PASTE_10X_MAP:
                    factor = COPY_PASTE_10X_MAP[clean_d]
                else:
                    try:
                        factor = parse_dilution_value(clean_d)
                    except ValueError:
                        factor = np.nan

            records.append({
                "Well": well_id,
                "Row": r_letter,
                "Col": col_idx,
                "Primer": primer,
                "Series": series,
                "Tag": tag.strip(),
                "Dilution_Factor": factor,
                "Is_NTC": is_ntc,
                "Is_Empty": False
            })

    return pd.DataFrame(records)


def clean_matrix_file(filepath: str, axis_name: str) -> pd.DataFrame:
    """Clean curve data matrices and drop duplicate columns."""
    if not filepath or not os.path.exists(filepath):
        return None
    df = pd.read_csv(filepath)
    unnamed = [c for c in df.columns if "unnamed" in str(c).lower()]
    if unnamed and len(df.columns) > 1:
        df = df.drop(columns=unnamed)

    new_cols = {}
    for c in df.columns:
        if axis_name.lower() in str(c).lower():
            new_cols[c] = axis_name
        else:
            new_cols[c] = normalize_well(c)

    df = df.rename(columns=new_cols)
    df = df.loc[:, ~df.columns.duplicated(keep="last")]
    return df


def calc_stats(x_vals, y_vals):
    if len(x_vals) < 3:
        return np.nan, np.nan, np.nan, np.nan
    slope, intercept, r_val, p_val, se = stats.linregress(x_vals, y_vals)
    eff = (10 ** (-1.0 / slope) - 1.0) * 100.0
    return slope, intercept, r_val ** 2, eff


# ----------------- Synthetic Demo Generator -----------------

def generate_synthetic_demo_data():
    """Generate artificial synthetic qPCR run for demonstration."""
    np.random.seed(42)
    rows_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
    records = []
    
    for r in rows_letters:
        for c in range(1, 13):
            well = f"{r}{c:02d}"
            primer = "TargetX" if c <= 6 else "TargetY"
            dil_idx = (c - 1) % 6
            if dil_idx < 5:
                dil_val = 10 ** dil_idx
                tag = f"{primer}_10x_{dil_val}"
                is_ntc = False
                cq = 22.0 + 3.32 * dil_idx + np.random.normal(0, 0.12)
                if primer == "TargetX" and dil_idx == 4:
                    cq = 34.2  # Plateau LOQ demonstration
            else:
                dil_val = np.nan
                tag = f"{primer}_10x_NTC"
                is_ntc = True
                cq = np.nan

            records.append({
                "Well": well, "Row": r, "Col": c, "Primer": primer, "Series": "10x",
                "Tag": tag, "Dilution_Factor": dil_val, "Is_NTC": is_ntc, "Cq": cq,
                "Is_Empty": False
            })

    merged_df = pd.DataFrame(records)

    cycles = np.arange(1, 41)
    amp_dict = {"Cycle": cycles}
    for _, row in merged_df.iterrows():
        w = row["Well"]
        if row["Is_NTC"]:
            amp_dict[w] = np.random.normal(5, 2, 40)
        else:
            cq_val = row["Cq"]
            amp_dict[w] = 3000.0 / (1.0 + np.exp(-(cycles - cq_val) / 1.4)) + np.random.normal(0, 8, 40)
    amp_df = pd.DataFrame(amp_dict)

    temps = np.linspace(60, 95, 71)
    melt_dict = {"Temperature": temps}
    for _, row in merged_df.iterrows():
        w = row["Well"]
        if row["Is_NTC"]:
            melt_dict[w] = np.random.normal(2, 1, len(temps))
        else:
            peak_t = 83.5 if row["Primer"] == "TargetX" else 81.8
            melt_dict[w] = 450.0 * np.exp(-((temps - peak_t) ** 2) / (2 * 1.5 ** 2)) + np.random.normal(0, 4, len(temps))
    melt_df = pd.DataFrame(melt_dict)

    return merged_df, amp_df, melt_df


# ----------------- Sidebar: Controls -----------------

with st.sidebar:
    st.markdown("## 🧬 qPCR primer validation")
    st.caption("Standard curve & MIQE quality control.")

    st.markdown("---")
    use_demo = st.checkbox("🧪 Load demo synthetic data", value=False, help="Try the tool with simulated data without uploading files.")

    active_data_dir = None
    run_name = "qPCR_Run"

    if not use_demo:
        # Default to ZIP upload first
        input_mode = st.radio("Input mode:", ["Upload ZIP archive", "Local folder"])
        
        if input_mode == "Upload ZIP archive":
            uploaded_zip = st.file_uploader("Upload .ZIP archive:", type=["zip"])
            if uploaded_zip is not None:
                temp_dir_obj = tempfile.TemporaryDirectory()
                zip_path = os.path.join(temp_dir_obj.name, uploaded_zip.name)
                with open(zip_path, "wb") as f:
                    f.write(uploaded_zip.getbuffer())
                with zipfile.ZipFile(zip_path, "r") as zf:
                    zf.extractall(temp_dir_obj.name)
                active_data_dir = temp_dir_obj.name
                run_name = os.path.splitext(uploaded_zip.name)[0]
        else:
            current_exec_dir = os.getcwd()
            detected_dirs = []
            for root, dirs, _ in os.walk(current_exec_dir):
                for d in dirs:
                    rel_p = os.path.relpath(os.path.join(root, d), current_exec_dir)
                    if not rel_p.startswith(".") and "env" not in rel_p and "git" not in rel_p:
                        detected_dirs.append(rel_p)

            selected_sub = st.selectbox(
                "Quick folder pick (optional):", 
                ["Current directory"] + sorted(detected_dirs)[:15]
            )
            
            initial_path = current_exec_dir if selected_sub == "Current directory" else os.path.join(current_exec_dir, selected_sub)
            local_path = st.text_input("Folder path:", value=initial_path)
            
            if os.path.exists(local_path):
                active_data_dir = local_path
                run_name = os.path.basename(os.path.abspath(local_path))
            else:
                st.error("Folder not found.")

        st.markdown("---")
        st.info("📋 **Expected files:**\n- `*Quantification Cq Results.csv`\n- `*Quantification Amplification Results*.csv`\n- `*Melt Curve Derivative Results*.csv`\n- `plate.xlsx` (or layout file)")


# ----------------- Data Loading -----------------

if use_demo:
    merged_df, amp_df, melt_df = generate_synthetic_demo_data()
    run_name = "Synthetic_Demo_Experiment"
    layout_name = "Simulated_Plate_96.xlsx"
else:
    if not active_data_dir:
        st.info("👋 Upload a ZIP archive or specify a folder path in the sidebar to begin analysis.")
        st.stop()

    found = find_files_recursively(active_data_dir)

    if not found["cq"]:
        st.error("❌ Could not find `*Quantification Cq Results.csv` in this folder.")
        st.stop()

    if not found["layout"]:
        st.error("❌ Plate layout file not found! Please place an Excel file named `plate.xlsx` or `layout.xlsx` in the folder.")
        st.stop()

    with st.spinner("Processing thermocycler dataset..."):
        cq_df = pd.read_csv(found["cq"])
        cq_df["Well"] = cq_df["Well"].apply(normalize_well)
        cq_df["Cq"] = pd.to_numeric(cq_df["Cq"], errors="coerce")

        amp_df = clean_matrix_file(found["amp"], "Cycle")
        melt_df = clean_matrix_file(found["melt"], "Temperature")
        layout_df = parse_plate_file(found["layout"])

        if layout_df.empty:
            st.error("Failed to parse 8x12 plate layout. Please ensure rows are labeled A through H.")
            st.stop()

        bad_tags = layout_df[~layout_df["Is_Empty"] & ~layout_df["Is_NTC"] & layout_df["Dilution_Factor"].isna()]
        if not bad_tags.empty:
            st.warning(
                f"⚠️ {len(bad_tags)} well(s) have a dilution that could not be read and are excluded "
                f"(expected `Primer_Series_Dilution`): {', '.join(sorted(bad_tags['Tag'].unique()))}"
            )

        cq_df = cq_df.drop_duplicates(subset="Well", keep="first")
        merged_df = pd.merge(layout_df, cq_df[["Well", "Cq"]], on="Well", how="left")
        layout_name = os.path.basename(found["layout"])


# ==============================================================================
#                      SINGLE SCROLLABLE DASHBOARD
# ==============================================================================

st.title(f"🧬 qPCR primer validation: {run_name}")
st.caption(f"Layout source: `{layout_name}` | Rows: 8 (A-H) | Columns: 12 (1-12)")

# Filter out empty wells from all downstream statistics
valid_df = merged_df[~merged_df["Is_Empty"] & (merged_df["Primer"] != "Empty")].copy()
# Wells whose dilution could not be read were already reported above; keep them out of the analysis
valid_df = valid_df[valid_df["Is_NTC"] | valid_df["Dilution_Factor"].notna()]

if valid_df.empty:
    st.warning("⚠️ No valid primer targets detected on the plate.")
    st.stop()

conditions = valid_df[["Primer", "Series"]].drop_duplicates().sort_values(by=["Primer", "Series"])

# ----------------- ANALYSIS: every primer, honouring the well checkboxes -----------------

def is_optimal(eff, r2) -> bool:
    return not (np.isnan(eff) or np.isnan(r2)) and 90 <= eff <= 110 and r2 >= 0.98


def status_of(eff, r2) -> str:
    if np.isnan(eff) or np.isnan(r2):
        return "NO DATA"
    return "OPTIMAL" if is_optimal(eff, r2) else "SUBOPTIMAL"


STATUS_ICON = {"OPTIMAL": "🟢 Optimal", "SUBOPTIMAL": "🔴 Suboptimal", "NO DATA": "⚪ No data"}


def fit_points(wells: pd.DataFrame):
    """Mean Cq per dilution, then regression: (grouped, slope, intercept, r2, eff)."""
    grp = wells.groupby("Dilution_Factor")["Cq"].agg(["mean", "std", "count"]).reset_index()
    if len(grp) >= 3:
        return (grp, *calc_stats(-np.log10(grp["Dilution_Factor"].values), grp["mean"].values))
    return grp, np.nan, np.nan, np.nan, np.nan


def well_key(well: str) -> str:
    return f"w_chk_{run_name}_{well}"


def analyze_condition(primer: str, series: str) -> dict:
    sub = valid_df[(valid_df["Primer"] == primer) & (valid_df["Series"] == series)]
    ntc = sub[sub["Is_NTC"]].sort_values("Well")
    std = sub[~sub["Is_NTC"]].dropna(subset=["Cq"]).sort_values(["Dilution_Factor", "Well"]).copy()
    n_no_cq = int((~sub["Is_NTC"] & sub["Cq"].isna()).sum())

    # All wells included
    grp_raw, _, _, r2_raw, eff_raw = fit_points(std)

    # Smart LOQ advisor: does the curve pass without the most dilute point?
    rec_factor, advice = None, ""
    if not np.isnan(eff_raw) and not is_optimal(eff_raw, r2_raw):
        trimmed = grp_raw.iloc[:-1]
        if len(trimmed) >= 3:
            _, _, r2_t, eff_t = calc_stats(-np.log10(trimmed["Dilution_Factor"].values), trimmed["mean"].values)
            if is_optimal(eff_t, r2_t):
                rec_factor = grp_raw.iloc[-1]["Dilution_Factor"]
                advice = (f"Smart advice: dilution {format_dilution(rec_factor)} reached the limit of quantification (plateau). "
                          f"Excluding it gives Eff {eff_t:.1f}% and R² {r2_t:.4f}.")

    # Current selection: checkbox state, defaulting to the advisor's recommendation
    std["Active"] = [
        st.session_state.get(well_key(w), not (rec_factor is not None and f == rec_factor))
        for w, f in zip(std["Well"], std["Dilution_Factor"])
    ]
    std["Active"] = std["Active"].astype(bool)  # an empty list would give object dtype and break masking
    grp_sel, slope, intercept, r2, eff = fit_points(std[std["Active"]])

    ntc_min = ntc["Cq"].min()  # worst replicate; a mean would hide contamination
    if pd.isna(ntc_min):
        ntc_label = "Clean" if not ntc.empty else "No NTC"
    else:
        ntc_label = f"Pass ({ntc_min:.1f})" if ntc_min >= 35 else f"High ({ntc_min:.1f})"

    return {
        "primer": primer, "series": series, "std": std, "ntc": ntc, "n_no_cq": n_no_cq,
        "eff_raw": eff_raw, "r2_raw": r2_raw, "grp_sel": grp_sel,
        "slope": slope, "intercept": intercept, "r2": r2, "eff": eff,
        "status": status_of(eff, r2), "ntc_label": ntc_label,
        "rec_factor": rec_factor, "advice": advice,
        "recommendation": f"Exclude {format_dilution(rec_factor)} (plateau/LOQ)" if rec_factor is not None else "Keep as-is",
        "n_active": int(std["Active"].sum()), "n_total": len(std),
        "excluded_wells": list(std.loc[~std["Active"], "Well"]),
    }


results = [analyze_condition(r["Primer"], r["Series"]) for _, r in conditions.iterrows()]

# ----------------- SECTION 1: EXECUTIVE SUMMARY -----------------
st.markdown("### 📊 Executive summary")
st.caption("All primers are analysed at once. Wells you untick below are reflected here, in the PDF and in the Excel report.")

sum_table = pd.DataFrame([{
    "Primer": r["primer"],
    "Series": r["series"],
    "Raw Eff (%)": f"{r['eff_raw']:.1f}%" if not np.isnan(r["eff_raw"]) else "No Amp",
    "Raw R²": f"{r['r2_raw']:.4f}" if not np.isnan(r["r2_raw"]) else "N/A",
    "Eff (%)": f"{r['eff']:.1f}%" if not np.isnan(r["eff"]) else "No Amp",
    "R²": f"{r['r2']:.4f}" if not np.isnan(r["r2"]) else "N/A",
    "Slope": f"{r['slope']:.3f}" if not np.isnan(r["slope"]) else "N/A",
    "Wells used": f"{r['n_active']}/{r['n_total']}",
    "Smart recommendation": r["recommendation"],
    "Status": STATUS_ICON[r["status"]],
    "NTC QC": r["ntc_label"],
} for r in results])
st.dataframe(sum_table, width="stretch", hide_index=True)

# Downloads: PDF (built on demand, it takes a few seconds) and Excel
selection_sig = (run_name, tuple((w, bool(a)) for r in results for w, a in zip(r["std"]["Well"], r["std"]["Active"])))
export_name = f"{run_name}_qPCR"
col_pdf, col_pdf_dl, col_xls = st.columns(3)

with col_pdf:
    if st.button("📄 Build PDF report", width="stretch"):
        with st.spinner("Rendering the PDF (plots for every primer)..."):
            st.session_state["pdf_report"] = (
                selection_sig,
                build_pdf_report(run_name, layout_name, merged_df, results, amp_df, melt_df),
            )

pdf_state = st.session_state.get("pdf_report")
with col_pdf_dl:
    if pdf_state and pdf_state[0] == selection_sig:
        st.download_button(
            "⬇️ Download PDF", data=pdf_state[1], file_name=f"{export_name}_report.pdf",
            mime="application/pdf", width="stretch",
        )
    elif pdf_state:
        st.caption("Selection changed since the last build. Build the PDF again.")

included = {w for r in results for w in r["std"].loc[r["std"]["Active"], "Well"]}
wells_out = merged_df[["Well", "Row", "Col", "Primer", "Series", "Tag", "Cq", "Is_NTC", "Is_Empty"]].copy()
wells_out["Included_in_fit"] = [
    "" if (ntc or emp) else ("yes" if w in included else "no")
    for w, ntc, emp in zip(wells_out["Well"], wells_out["Is_NTC"], wells_out["Is_Empty"])
]

excel_buf = io.BytesIO()
with pd.ExcelWriter(excel_buf, engine="openpyxl") as writer:
    sum_table.to_excel(writer, sheet_name="Executive_Summary", index=False)

    # Sheet 2: 96-well plate map (guaranteed 8x12 grid via reindex)
    ws_plate = writer.book.create_sheet(title="Plate_96_Map")
    rows_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
    cols_numbers = list(range(1, 13))

    plate_cq_grid = merged_df.pivot(index="Row", columns="Col", values="Cq").reindex(index=rows_letters, columns=cols_numbers)
    plate_tag_grid = merged_df.pivot(index="Row", columns="Col", values="Tag").reindex(index=rows_letters, columns=cols_numbers)

    hdr_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sec_font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    border = Border(left=Side(style="thin", color="D3D3D3"), right=Side(style="thin", color="D3D3D3"),
                    top=Side(style="thin", color="D3D3D3"), bottom=Side(style="thin", color="D3D3D3"))

    ws_plate["A1"] = "1. Plate layout map (primers & dilutions):"
    ws_plate["A1"].font = sec_font
    for c in range(1, 13):
        cell = ws_plate.cell(row=3, column=c + 1, value=c)
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = Alignment(horizontal="center")
    for r_idx, r_let in enumerate(rows_letters, start=4):
        ws_plate.cell(row=r_idx, column=1, value=r_let).font = Font(bold=True)
        for c_idx in range(1, 13):
            tag_val = plate_tag_grid.loc[r_let, c_idx]
            cell = ws_plate.cell(row=r_idx, column=c_idx + 1, value="" if pd.isna(tag_val) or tag_val == "Empty" else str(tag_val))
            cell.border = border; cell.alignment = Alignment(horizontal="center")

    ws_plate["A15"] = "2. Bio-Rad instrument Cq values:"
    ws_plate["A15"].font = sec_font
    for c in range(1, 13):
        cell = ws_plate.cell(row=17, column=c + 1, value=c)
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = Alignment(horizontal="center")
    for r_idx, r_let in enumerate(rows_letters, start=18):
        ws_plate.cell(row=r_idx, column=1, value=r_let).font = Font(bold=True)
        for c_idx in range(1, 13):
            val = plate_cq_grid.loc[r_let, c_idx]
            tag_val = plate_tag_grid.loc[r_let, c_idx]
            cell = ws_plate.cell(row=r_idx, column=c_idx + 1)
            cell.border = border; cell.alignment = Alignment(horizontal="center")
            if pd.isna(val):
                if not (pd.isna(tag_val) or tag_val == "Empty"):
                    cell.value = "No Cq"
                    cell.font = Font(color="808080", italic=True)
            else:
                cell.value = round(val, 2)
                cell.number_format = "0.00"

    wells_out.to_excel(writer, sheet_name="Raw_Wells", index=False)

    for ws in [writer.sheets["Executive_Summary"], writer.sheets["Plate_96_Map"], writer.sheets["Raw_Wells"]]:
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 12)

with col_xls:
    st.download_button(
        "📊 Download Excel report", data=excel_buf.getvalue(), file_name=f"{export_name}_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch",
    )

st.divider()

# ----------------- SECTION 2: 96-WELL HEATMAP WITH CONTRAST NUMBERS -----------------
st.markdown("### 🧪 96-well plate Cq heatmap")
st.caption("Each well displays its exact Cq value. Contrast colors auto-adapt (white text on dark blue, dark text on light blue).")

rows_order = ["H", "G", "F", "E", "D", "C", "B", "A"]
cols_order = list(range(1, 13))

cq_matrix = np.full((8, 12), np.nan)
hover_matrix = np.empty((8, 12), dtype=object)
annotations = []

for r_idx, r_let in enumerate(rows_order):
    for c_idx, c_num in enumerate(cols_order):
        well_id = f"{r_let}{c_num:02d}"
        match = merged_df[merged_df["Well"] == well_id]
        if not match.empty:
            val = match.iloc[0]["Cq"]
            tag = match.iloc[0]["Tag"]
            is_emp = match.iloc[0].get("Is_Empty", False)
            cq_matrix[r_idx, c_idx] = val
            
            if is_emp or tag == "Empty":
                txt = "—"
                font_col = "#AAAAAA"
                hover_matrix[r_idx, c_idx] = f"Well: <b>{well_id}</b><br><i>Unused (Empty)</i>"
            elif pd.isna(val):
                txt = "—"
                font_col = "#888888"
                hover_matrix[r_idx, c_idx] = f"Well: <b>{well_id}</b><br>Sample: {tag}<br><i>No amplification</i>"
            else:
                txt = f"{val:.1f}"
                font_col = "#FFFFFF" if val < 27 else "#111111"
                hover_matrix[r_idx, c_idx] = f"Well: <b>{well_id}</b><br>Sample: {tag}<br>Cq: <b>{val:.2f}</b>"

            annotations.append(dict(
                x=str(c_num),
                y=r_let,
                text=txt,
                showarrow=False,
                font=dict(color=font_col, size=11, family="Arial")
            ))

fig_heat = go.Figure(data=go.Heatmap(
    z=cq_matrix,
    x=[str(c) for c in cols_order],
    y=rows_order,
    hovertext=hover_matrix,
    hoverinfo="text",
    colorscale="Blues_r",
    colorbar=dict(title="Cq"),
    xgap=4,
    ygap=4
))

fig_heat.update_layout(
    annotations=annotations,
    height=420,
    margin=dict(l=40, r=40, t=10, b=40)
)
fig_heat.update_xaxes(tickmode="linear", tick0=1, dtick=1, title="Column (1 - 12)")
fig_heat.update_yaxes(tickmode="linear", dtick=1, title="Row (A - H)")

st.plotly_chart(fig_heat, width="stretch")

st.divider()

# ----------------- SECTION 3: ALL PRIMERS, ONE AFTER ANOTHER -----------------
st.markdown("### 🔍 Primer details & well manager")
st.caption("Every primer is shown below. Untick a replicate to exclude it; numbers, plots, PDF and Excel update.")

for res in results:
    primer, series = res["primer"], res["series"]
    std = res["std"]
    cond_data = valid_df[(valid_df["Primer"] == primer) & (valid_df["Series"] == series)]
    active_wells = set(std.loc[std["Active"], "Well"])
    k = f"{primer}_{series}"

    with st.container(border=True):
        st.markdown(f"#### {primer} ({series})")

        if std.empty:
            st.error("No Cq values for this primer: no amplification detected in any sample well.")
        elif res["rec_factor"] is not None:
            st.warning("💡 " + res["advice"] + " Those wells are unticked by default; tick them back if you disagree.")
        elif res["status"] == "OPTIMAL":
            st.success("✅ Standard curve satisfies the guidelines (Eff 90-110%, R² >= 0.98).")
        elif res["status"] == "NO DATA":
            st.info("Fewer than 3 dilution points selected: no standard curve can be fitted.")
        else:
            st.error("Standard curve is outside the guidelines (Eff 90-110%, R² >= 0.98) and removing the last dilution does not fix it.")

        if res["n_no_cq"]:
            st.caption(f"⚠️ {res['n_no_cq']} sample well(s) have no Cq.")

        with st.expander("🛠️ Interactive well manager (toggle replicates)", expanded=True):
            st.write("Uncheck any replicate or dilution point to observe real-time recalculation:")
            chk_cols = st.columns(5)
            for i, (_, w_row) in enumerate(std.iterrows()):
                chk_cols[i % 5].checkbox(
                    f"{w_row['Well']} ({format_dilution(w_row['Dilution_Factor'])}): Cq={w_row['Cq']:.2f}",
                    value=bool(w_row["Active"]),
                    key=well_key(w_row["Well"]),
                )

        eff_delta = (f"{res['eff'] - res['eff_raw']:+.1f}% vs baseline"
                     if not np.isnan(res["eff_raw"]) and not np.isnan(res["eff"]) and res["n_active"] != res["n_total"] else None)
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("PCR efficiency", f"{res['eff']:.2f}%" if not np.isnan(res["eff"]) else "N/A", delta=eff_delta)
        k2.metric("R-squared (R²)", f"{res['r2']:.4f}" if not np.isnan(res["r2"]) else "N/A")
        k3.metric("Slope", f"{res['slope']:.3f}" if not np.isnan(res["slope"]) else "N/A")
        k4.metric("Live status", STATUS_ICON[res["status"]])

        c_p1, c_p2, c_p3 = st.columns(3)
        for col, df_, xcol, key, color, title, xt, yt in (
            (c_p1, amp_df, "Cycle", "amp", "#1f77b4", "Amplification curves", "Cycle", "RFU"),
            (c_p2, melt_df, "Temperature", "melt", "#2ca02c", "Melt peaks (-d(RFU)/dT)", "Temp (°C)", "-d(RFU)/dT"),
        ):
            with col:
                fig = go.Figure()
                if df_ is not None and xcol in df_.columns:
                    for _, w_r in cond_data.iterrows():
                        w = w_r["Well"]
                        if w in df_.columns:
                            is_ntc, is_active = w_r["Is_NTC"], w in active_wells
                            name = "NTC" if is_ntc else f"{format_dilution(w_r['Dilution_Factor'])} ({w})"
                            fig.add_trace(go.Scatter(
                                x=df_[xcol].values, y=df_[w].values, mode="lines", name=name,
                                line=dict(color="red" if is_ntc else (color if is_active else "#dddddd"),
                                          dash="dash" if is_ntc else "solid",
                                          width=1.5 if is_active or is_ntc else 0.8)))
                fig.update_layout(title=title, xaxis_title=xt, yaxis_title=yt, height=380, showlegend=False)
                st.plotly_chart(fig, width="stretch", key=f"{key}_{k}")

        with c_p3:
            f_std = go.Figure()
            g = res["grp_sel"]
            if len(g) >= 3:
                act_x = -np.log10(g["Dilution_Factor"].values)
                act_y = g["mean"].values
                act_sd = g["std"].fillna(0).values
                f_std.add_trace(go.Scatter(
                    x=act_x, y=act_y, mode="markers",
                    error_y=dict(type="data", array=act_sd, visible=True),
                    marker=dict(size=10, color="#1F4E79"),
                    hovertext=[f"Dilution: {format_dilution(d)}<br>Mean Cq: {m:.2f} ± {s:.2f}<br>Active wells: {int(c)}"
                               for d, m, s, c in zip(g["Dilution_Factor"], act_y, act_sd, g["count"])],
                    hoverinfo="text", name="Active points"))
                x_line = np.linspace(min(act_x) - 0.2, max(act_x) + 0.2, 50)
                f_std.add_trace(go.Scatter(x=x_line, y=res["slope"] * x_line + res["intercept"], mode="lines",
                                           line=dict(color="#C00000", width=2), name="Fit"))
            std_title = ("Standard curve (need ≥ 3 points)" if np.isnan(res["eff"])
                         else f"Standard curve (Eff={res['eff']:.1f}%, R²={res['r2']:.3f})")
            f_std.update_layout(title=std_title, xaxis_title="Log10 rel conc", yaxis_title="Cq", height=380, showlegend=False)
            st.plotly_chart(f_std, width="stretch", key=f"std_{k}")
